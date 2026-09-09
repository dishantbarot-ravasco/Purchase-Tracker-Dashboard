"""
apps/services/parsers/stock.py — parses HRS's RAW MATERIAL STOCK.xlsx,
'Stock' sheet, into a flat list of row dicts ready to upsert into
HRSRMLot.

Exact header (row 6; data from row 7), verified against the live file this
session:
A S.No. | B Description | C SAP ITEM CODE | D Category | E Sub Category |
F UOM | G Opening Stock | H REC | I ISSUE | J Today Stock | K Basic Rate |
L Value | M Rec. DT. | N No of Days | O (unlabeled - Party Name) |
P (unlabeled - location tag, e.g. 'HRS'/'RTP-1')

REC, ISSUE, Today Stock, Value, and No of Days are formula cells in the
source sheet (REC/ISSUE pull from that file's own Receipt/Issue tabs by
fixed row position - see project history). Reading with data_only=True
gives the cached computed value, which is what we want here; we are not
trying to re-derive or validate those formulas, just capture the number
the sheet is currently showing. Note that `received` (REC) reads 0 for
nearly every real lot in the live file - it appears to clear once a lot's
receipt is allocated rather than holding a running total, which is why
PO<->MIR<->Stock matching never compares stock quantity, only rate (see
apps/services/matching.py).

Column O ("Party Name") is one row per (material, vendor) lot, not one row
per material - the same material can appear many times from different
vendors at different rates. Its value carries a city suffix in the live
file (e.g. "Rubamin Private Limited - Vadodara") that MIR/PO data never
does, which is why vendor-name matching against this column uses
containment rather than exact equality (see common.py's normalize_vendor()
and matching.py's _vendor_matches()).
"""

import io
from dataclasses import dataclass

import openpyxl
from openpyxl.utils import column_index_from_string

from apps.services.parsers.common import stream_rows, to_date, to_decimal, to_str

# ── Column/row layout constants ─────────────────────────────────────────────

SHEET_NAME = "Stock"
HEADER_ROW = 6
DATA_START_ROW = 7

# Only A-N are header-checked - O and P are genuinely unlabeled in the
# source file (confirmed this session), so there's no header text to
# validate them against.
EXPECTED_HEADERS = {
    "A": "S.No.", "B": "Description", "C": "SAP ITEM CODE", "D": "Category",
    "E": "Sub Category", "F": "UOM", "G": "Opening\nStock", "H": "REC", "I": "ISSUE",
    "J": "Today\nStock", "K": "Basic Rate", "L": "Value", "M": "Rec. DT.", "N": "No of Days",
}


# ── Parsed-row shape ─────────────────────────────────────────────────────────

@dataclass
class ParsedStockLot:
    sr_no: object
    description: str
    sap_item_code: str
    category: str
    sub_category: str
    uom: str
    opening_stock: object
    received: object
    issued: object
    todays_stock: object
    basic_rate: object
    value: object
    received_date: object
    no_of_days: object
    party_name: str
    location_tag: str
    source_row_ref: str


class HeaderMismatch(Exception):
    pass


def _strip(v):
    return (v or "").strip() if isinstance(v, str) else v


# ── Public entry point ───────────────────────────────────────────────────────

def parse_stock_xlsx(file_bytes: bytes) -> list[ParsedStockLot]:
    """Reads the 'Stock' sheet and returns one ParsedStockLot per lot row.
    Raises HeaderMismatch immediately if the sheet name or any checked
    header cell doesn't match what this parser was built against."""
    # read_only=True (added 2026-09-08, real production OOM found on Render's
    # Starter-tier worker - 512MB): default mode loads the ENTIRE workbook
    # object graph into memory, including pivot table caches this app never
    # reads at all (confirmed via openpyxl's own "invalid dependency
    # definitions" warnings on every pivotCacheDefinition part of the real
    # file) - read_only mode streams instead, and explicitly skips pivot
    # tables/charts/styles it doesn't need. The data loop below streams via
    # common.stream_rows() rather than per-cell random access (ws['A5'] etc.)
    # or ws.max_row/max_column - see that function's own docstring for why
    # (both can be None in read_only mode when the workbook's <dimension> tag
    # is broken, and random access on a read_only worksheet is O(n) per call,
    # not O(1) - both confirmed against real production data/files).
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise HeaderMismatch(f"Expected sheet {SHEET_NAME!r}, found sheets: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    for col, expected in EXPECTED_HEADERS.items():
        actual = _strip(ws[f"{col}{HEADER_ROW}"].value)
        if actual != expected:
            raise HeaderMismatch(f"Column {col}{HEADER_ROW}: expected {expected!r}, found {actual!r}")

    lots = []
    max_col = column_index_from_string("P")
    for r, c in stream_rows(ws, DATA_START_ROW, max_col):
        description = to_str(c["B"].value)
        if not description:
            continue

        sr_no_raw = c["A"].value
        sr_no = int(sr_no_raw) if isinstance(sr_no_raw, (int, float)) else None

        no_of_days_raw = to_decimal(c["N"].value)
        no_of_days = int(no_of_days_raw) if no_of_days_raw is not None else None

        lots.append(
            ParsedStockLot(
                sr_no=sr_no,
                description=description,
                sap_item_code=to_str(c["C"].value),
                category=to_str(c["D"].value),
                sub_category=to_str(c["E"].value),
                uom=to_str(c["F"].value),
                opening_stock=to_decimal(c["G"].value) or 0,
                received=to_decimal(c["H"].value) or 0,
                issued=to_decimal(c["I"].value) or 0,
                todays_stock=to_decimal(c["J"].value) or 0,
                basic_rate=to_decimal(c["K"].value),
                value=to_decimal(c["L"].value),
                received_date=to_date(c["M"].value),
                no_of_days=no_of_days,
                party_name=to_str(c["O"].value),
                location_tag=to_str(c["P"].value),
                source_row_ref=str(r),
            )
        )
    return lots
