"""
apps/services/parsers/mir.py - parses HRS's MIR FILE 2026-2027.xlsx,
'RAW MATERIAL' sheet, into a flat list of row dicts ready to upsert into
HRSMIREntry.

Exact header (row 6; data from row 7), verified against the live file this
session:
A Month | B MIR. No. | C Date | D SAP GRN No. | E Party Name | F State |
G Invoice No. | H Date | I Park Invoice No. | J Purchase Order. No. |
K Purchase Order. Date | L SAP P.O. No. | M SAP P.O. DATE |
N Material Description | O Qty | P UOM | Q Rate | R Net | S Discount Rate |
T Discount Amt. | U Others | V Taxable Value | W Tax rate | X IGST |
Y CGST Rate | Z CGST AMT | AA SGST RATE | AB SGST AMT |
AC Other exclu. gst | AD Total Amount | AE TCS RATE | AF TCS Amt. |
AG Inv. Final Value | AH Plant | AI Dept. Uses | AJ Category material |
AK Remarks

Column J ("Purchase Order. No.") is the field the docstring on
HRSMIREntry.po_number_raw warns about - confirmed unreliable by audit
(~30% blank, ~25% non-standard format) and never used as a sole join key.

Column C ("Date") is formatted mm/dd/yyyy and Excel has committed
day/month-transposed values to 177 of its 497 real date cells - see
_repair_mir_date() for the measurement and for why the repair keys on the
MIR number's own month.

Columns A ("Month") and B ("MIR. No.") have the identical Excel
autoconvert quirk documented in apps/services/parsers/achhad_mir.py's module
docstring - confirmed against the live file this session (previously
undocumented and unhandled here, unlike achhad_mir.py): typed text like
"April-26" / "01/04" gets silently turned into a real datetime cell by
Excel (number_format 'mmmm-d' / 'mm/dd' respectively) rather than staying
literal text. Before this fix, plain to_str() on these columns returned
Python's default str(datetime) repr (e.g. "2026-04-26 00:00:00") instead
of the originally-typed label - confirmed 100% of HRS's `month` values and
16.3% of `mir_no` values (day-of-month <=12, same as achhad_mir.py's
finding) were corrupted this way in the live data before this fix.
_month_label()/_mir_no_label() below are a straight port of
achhad_mir.py's identically-named helpers - the reconstruction logic is
generic, not Achhad-specific, it just never got applied here.
"""

import datetime
import io
import re
from dataclasses import dataclass

import openpyxl
from openpyxl.utils import column_index_from_string

from apps.services.parsers.common import (
    repair_month_swapped_date,
    stream_rows,
    to_date,
    to_decimal,
    to_str,
)

# ── Column/row layout constants ─────────────────────────────────────────────

SHEET_NAME = "RAW MATERIAL"
HEADER_ROW = 6
DATA_START_ROW = 7

EXPECTED_HEADERS = {
    "A": "Month", "B": "MIR. No.", "C": "Date", "D": "SAP GRN No.", "E": "Party Name",
    "F": "State", "G": "Invoice No.", "H": "Date", "I": "Park Invoice No.",
    "J": "Purchase Order. No.", "K": "Purchase Order. Date", "L": "SAP P.O. No.",
    "M": "SAP P.O. DATE", "N": "Material Description", "O": "Qty", "P": "UOM", "Q": "Rate",
    "R": "Net", "S": "Discount Rate", "T": "Discount Amt.", "U": "Others",
    "V": "Taxable Value", "W": "Tax rate", "X": "IGST", "Y": "CGST Rate", "Z": "CGST AMT",
    "AA": "SGST RATE", "AB": "SGST AMT", "AC": "Other exclu. gst", "AD": "Total Amount",
    "AE": "TCS RATE", "AF": "TCS Amt.", "AG": "Inv. Final Value", "AH": "Plant",
    "AI": "Dept. Uses", "AJ": "Category material", "AK": "Remarks",
}


# ── Parsed-row shape ─────────────────────────────────────────────────────────

@dataclass
class ParsedMirEntry:
    month: str
    mir_no: str
    mir_date: object
    sap_grn_number: str
    po_number_raw: str
    party_name: str
    state: str
    invoice_no: str
    invoice_date: object
    material_description: str
    qty: object
    uom: str
    rate: object
    net: object
    discount_rate_pct: object
    discount_amt: object
    others: object
    taxable_value: object
    tax_rate_pct: object
    igst: object
    cgst_rate_pct: object
    cgst_amt: object
    sgst_rate_pct: object
    sgst_amt: object
    other_taxes_excl_gst: object
    total_amount: object
    tcs_rate_pct: object
    tcs_amt: object
    invoice_final_value: object
    plant_tag: str
    dept_use: str
    material_category: str
    remarks: str
    source_row_ref: str


class HeaderMismatch(Exception):
    pass


def _strip(v):
    return (v or "").strip() if isinstance(v, str) else v


# ── Excel-autoconvert workarounds (Month/MIR No. columns) ───────────────────

def _month_label(cell) -> str:
    """See module docstring - reconstructs the originally-typed 'Mon-YY'
    text from a cell Excel silently turned into a date, using the cell's
    own number_format to know whether it should read as abbreviated or
    full month name. Ported verbatim from achhad_mir.py's identically-named
    helper."""
    v = cell.value
    if not isinstance(v, datetime.datetime):
        return to_str(v)
    month_name = v.strftime("%B") if "mmmm" in (cell.number_format or "") else v.strftime("%b")
    return f"{month_name}-{v.day}"


def _mir_no_label(cell) -> str:
    """See module docstring - reconstructs the originally-typed 'MM/DD'
    MIR-number text from a cell Excel silently turned into a date. Ported
    verbatim from achhad_mir.py's identically-named helper."""
    v = cell.value
    if not isinstance(v, datetime.datetime):
        return to_str(v)
    return f"{v.month:02d}/{v.day:02d}"


# ── Transposed-date repair (DATE column, 2026-09-18) ────────────────────────
# The same defect vapi_mir.py has carried a repair for since 2026-09-12, found
# in HRS's own file this session and repaired the same way. See
# _repair_mir_date() below for the measurement.

_MIR_NO_MONTH_RE = re.compile(r"/\s*(\d{1,2})\s*$")


def _mir_no_month(mir_no: str) -> int | None:
    """The month encoded in a MIR number - '24/04' is April's 24th MIR.

    Safe to read off _mir_no_label()'s output rather than the raw cell: when
    Excel has turned the MIR number into a date, that helper rebuilds the
    label from the SAME month/day pair Excel stored, so serial and month come
    back in the order they were typed (see its docstring). A serial above 12
    is never converted at all and stays literal text."""
    match = _MIR_NO_MONTH_RE.search(mir_no or "")
    if not match:
        return None
    month = int(match.group(1))
    return month if 1 <= month <= 12 else None


def _repair_mir_date(value, mir_no: str):
    """Undoes a day/month transposition in the DATE column using the month
    the MIR number states independently.

    MEASURED ON THE LIVE HRS FILE (2026-09-18): 177 of 497 real date cells
    are stored by Excel itself with day and month swapped - MIR '01/04'
    (April) carries the datetime 2026-01-04, MIR '16/08' carries
    2026-12-08 - and **38 rows land in the future**, October to December
    2026. Column C is formatted mm/dd/yyyy, so a date typed "5-4-2026"
    meaning 5 April was read month-first and committed to the file as 4 May.
    openpyxl hands back a genuine datetime; only a second, independent
    record of the month can tell. It is getting worse, not better: 28% of
    April's rows against 59% of September's.

    WHY THIS IS A SAFETY NET AND NOT THE FIX. The real repair is the cell
    format, at the plant (the message in PO_MIR_Audit/ asks for exactly
    that). This is what keeps the damage out of the database while that
    happens, and what catches it silently if the format is ever set back.
    Once the source is correct this function stops firing on its own - it
    only ever rewrites a value that is *exactly* transposed against the MIR
    number, so a correct date is never touched.

    WHAT IT DOES NOT COVER. Only the DATE column. The same mm/dd/yyyy format
    sits on Invoice Date (39 future-dated), Purchase Order Date and SAP P.O.
    Date (4 each), but none of those has an independent month to check
    against - an invoice may legitimately be raised in a different month
    from the receipt, so there is no safe rule and guessing one would be
    worse than leaving them alone. Those need the plant-side format fix.

    See repair_month_swapped_date() for the three conditions under which a
    value is left alone - a legitimately cross-month row is never
    rewritten."""
    return repair_month_swapped_date(to_date(value), _mir_no_month(mir_no))


# ── Public entry point ───────────────────────────────────────────────────────

def parse_mir_xlsx(file_bytes: bytes) -> list[ParsedMirEntry]:
    """Reads the 'RAW MATERIAL' sheet and returns one ParsedMirEntry per data
    row. Raises HeaderMismatch immediately if the sheet name or any header
    cell doesn't match what this parser was built against, rather than
    silently reading misaligned columns."""
    # read_only=True - see parsers/stock.py's own comment on this same line
    # for why (a real production OOM on Render, caused by default-mode
    # loading pivot table caches this app never reads). The data loop below
    # streams via common.stream_rows() rather than per-cell random access -
    # see that function's own docstring for why (ws.max_row/max_column can be
    # None in read_only mode, and random access on a read_only worksheet is
    # O(n) per call, not O(1) - both confirmed against real production data).
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise HeaderMismatch(f"Expected sheet {SHEET_NAME!r}, found sheets: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    for col, expected in EXPECTED_HEADERS.items():
        actual = _strip(ws[f"{col}{HEADER_ROW}"].value)
        if actual != expected:
            raise HeaderMismatch(f"Column {col}{HEADER_ROW}: expected {expected!r}, found {actual!r}")

    entries = []
    max_col = column_index_from_string("AK")
    for r, c in stream_rows(ws, DATA_START_ROW, max_col):
        party_name = to_str(c["E"].value)
        if not party_name:
            continue  # a blank Party Name means an empty row - MIR has no other reliable "is this row used" signal

        # mir_no is read first: it is what _repair_mir_date() checks the DATE
        # column against (see that function).
        mir_no = _mir_no_label(c["B"])
        entries.append(
            ParsedMirEntry(
                month=_month_label(c["A"]),
                mir_no=mir_no,
                mir_date=_repair_mir_date(c["C"].value, mir_no),
                sap_grn_number=to_str(c["D"].value),
                po_number_raw=to_str(c["J"].value),  # unreliable join key, see module docstring - never used alone
                party_name=party_name,
                state=to_str(c["F"].value),
                invoice_no=to_str(c["G"].value),
                invoice_date=to_date(c["H"].value),
                material_description=to_str(c["N"].value),
                qty=to_decimal(c["O"].value),
                uom=to_str(c["P"].value),
                rate=to_decimal(c["Q"].value),
                net=to_decimal(c["R"].value),
                discount_rate_pct=to_decimal(c["S"].value),
                discount_amt=to_decimal(c["T"].value),
                others=to_decimal(c["U"].value),
                taxable_value=to_decimal(c["V"].value),
                tax_rate_pct=to_decimal(c["W"].value),
                igst=to_decimal(c["X"].value),
                cgst_rate_pct=to_decimal(c["Y"].value),
                cgst_amt=to_decimal(c["Z"].value),
                sgst_rate_pct=to_decimal(c["AA"].value),
                sgst_amt=to_decimal(c["AB"].value),
                other_taxes_excl_gst=to_decimal(c["AC"].value),
                total_amount=to_decimal(c["AD"].value),
                tcs_rate_pct=to_decimal(c["AE"].value),
                tcs_amt=to_decimal(c["AF"].value),
                invoice_final_value=to_decimal(c["AG"].value),
                plant_tag=to_str(c["AH"].value),
                dept_use=to_str(c["AI"].value),
                material_category=to_str(c["AJ"].value),
                remarks=to_str(c["AK"].value),
                source_row_ref=str(r),
            )
        )
    return entries
