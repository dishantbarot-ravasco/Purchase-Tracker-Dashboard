"""
apps/services/parsers/rodtep.py — parses one RoDTEP scrip ledger xlsx (e.g.
"RODTEP-JNPT-1.xlsx") from the "Purchase Orders HO/RODTEP SCRIPT LICENSE"
Drive folder into a flat list of row dicts ready to upsert into
RodtepScrollEntry.

Genuinely different from every other Drive source this app parses: there is
no single fixed file title to search for (see google_client.py's
list_files_in_folder(), added for this) - a new "RODTEP-JNPT-<N>.xlsx" file
appears each time a new RoDTEP Script Number is issued, each holding that
one script's own Shipping-Bill-level breakdown. sync_rodtep.py lists the
whole folder and calls this parser once per file found.

Exact header, confirmed against two real files this session:
Sr No | Script No | Date | SB Number | SB Date | Scroll Number | Scroll
Date | Scroll Type | Location | Sanctioned Amount

One row per Shipping Bill (SB Number) that contributed credit to that
file's own Script No - `Sanctioned Amount` is the real RoDTEP credit (INR)
earned from that shipment. Confirmed real column meanings: `Script No` is
the ICEGATE scrip identifier (same value repeated down the whole file -
one script per file); `SB Number`/`SB Date` are the EXPORT Shipping Bill
that earned the credit (NOT an import document - nothing in this file
itself references an import PO/BOE at all; see RodtepUsage for how the
import side is tracked, manually, since Drive has no structured link
between a script and which import it was later used against - a copy of
this same file gets dropped into the relevant import PO's own Drive folder
for reference, but that's a folder-placement fact, not a data column).

**The header row's position is NOT consistent between real files** -
confirmed directly: one live file has the header on row 1 (data from row
2), another has a fully blank row 1 then the header on row 2 (data from row
3) - apparently whichever row a person last cleared/inserted in Excel
before saving. This parser scans the first few rows for the expected
header instead of assuming a fixed row number, and raises HeaderMismatch
only if it isn't found within that scan window - the same "fail loudly
rather than silently misread" principle every other parser in this app
already follows, just with a small tolerance for this file's own real
sloppiness instead of a fixed HEADER_ROW constant.
"""

import io
from dataclasses import dataclass

import openpyxl

from apps.services.parsers.common import stream_rows, to_code_str, to_date, to_decimal, to_str

# ── Column layout ────────────────────────────────────────────────────────────

EXPECTED_HEADERS = [
    "Sr No", "Script No", "Date", "SB Number", "SB Date", "Scroll Number",
    "Scroll Date", "Scroll Type", "Location", "Sanctioned Amount",
]
_MAX_HEADER_SCAN_ROW = 5  # give up looking for the header past this row


# ── Parsed-row shape ─────────────────────────────────────────────────────────

@dataclass
class ParsedRodtepEntry:
    script_no: str
    script_date: object
    sb_number: str
    sb_date: object
    scroll_number: str
    scroll_date: object
    scroll_type: str
    location: str
    sanctioned_amount: object
    source_row_ref: str


class HeaderMismatch(Exception):
    pass


def _strip(v):
    return (v or "").strip() if isinstance(v, str) else v


def _find_header_row(ws) -> int:
    """Scans rows 1..._MAX_HEADER_SCAN_ROW for a row whose cell values match
    EXPECTED_HEADERS exactly. Returns that row number. Raises HeaderMismatch
    if none of the scanned rows match."""
    for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=_MAX_HEADER_SCAN_ROW), start=1):
        values = [_strip(cell.value) for cell in row[: len(EXPECTED_HEADERS)]]
        if values == EXPECTED_HEADERS:
            return row_num
    raise HeaderMismatch(
        f"Expected header {EXPECTED_HEADERS} not found in rows 1-{_MAX_HEADER_SCAN_ROW}."
    )


# ── Public entry point ───────────────────────────────────────────────────────

def parse_rodtep_xlsx(file_bytes: bytes) -> list[ParsedRodtepEntry]:
    """Reads one RoDTEP scrip ledger workbook (single sheet) and returns one
    ParsedRodtepEntry per Shipping Bill row. Raises HeaderMismatch if the
    expected header isn't found within the first few rows."""
    # read_only=True - see stock.py's own comment on this same line for why
    # (a real production OOM on Render). Streams via common.stream_rows()
    # rather than per-cell random access - see that function's own docstring
    # for why (both ws.max_row/max_column being None, and O(n) per random
    # cell access, are real, already-hit production bugs in this app).
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row = _find_header_row(ws)
    data_start_row = header_row + 1
    max_col = len(EXPECTED_HEADERS)

    entries = []
    for r, c in stream_rows(ws, data_start_row, max_col):
        script_no = to_code_str(c[2].value)
        if not script_no:
            continue  # a blank Script No means an empty trailing row - no other reliable "is this row used" signal

        entries.append(
            ParsedRodtepEntry(
                script_no=script_no,
                script_date=to_date(c[3].value),
                sb_number=to_code_str(c[4].value),
                sb_date=to_date(c[5].value),
                scroll_number=to_code_str(c[6].value),
                scroll_date=to_date(c[7].value),
                scroll_type=to_str(c[8].value),
                location=to_str(c[9].value),
                sanctioned_amount=to_decimal(c[10].value) or 0,
                source_row_ref=str(r),
            )
        )
    return entries
