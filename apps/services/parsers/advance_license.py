"""
apps/services/parsers/advance_license.py — parses the Advance License Data
Google Sheet (one fixed file, referenced directly by Drive file id -
settings.ADVANCE_LICENSE_FILE_ID - maintained by hand by the project owner,
see google_client.download_spreadsheet_bytes_by_id()) into a list of
ParsedAdvanceLicense, each carrying its own ParsedAdvanceLicenseMaterial
rows.

Column layout - the exact header the project owner's own template uses
(confirmed against the real "Advance License.csv.xlsx" template file this
was built from):

License Number | License Issue Date | IEC | CIF Value Authorized (INR) |
FOB Value Export Target (INR) | Export Product Description |
Export Obligation Period End Date (Export Validity) | Import Validity End Date |
Status | Input Material Description | Input ITC(HS) Code |
Qty Authorized (this material, KGS) | CIF Value Authorized (this material, INR) |
Duty Saved % (this material) | BOE Number | BOE Date | Import PO Number |
Qty Imported (this BOE, KGS) | Value Imported / Duty Saved (this BOE, INR)

One row per (license, input material, usage instance) - the first 9 columns
(license-level) repeat identically across every row of the same License
Number; the remaining 10 columns are per-material, with the last 5 (usage)
blank until that specific material has actually been imported against.

Same "header position isn't guaranteed" tolerance as rodtep.py's own parser
- this file is hand-maintained in a spreadsheet editor, same real-world
sloppiness risk.
"""

import io
from dataclasses import dataclass, field

import openpyxl

from apps.services.parsers.common import stream_rows, to_code_str, to_date, to_decimal, to_str

# ── Column layout ────────────────────────────────────────────────────────────

EXPECTED_HEADERS = [
    "License Number", "License Issue Date", "IEC", "CIF Value Authorized (INR)",
    "FOB Value Export Target (INR)", "Export Product Description",
    "Export Obligation Period End Date (Export Validity)", "Import Validity End Date", "Status",
    "Input Material Description", "Input ITC(HS) Code", "Qty Authorized (this material, KGS)",
    "CIF Value Authorized (this material, INR)", "Duty Saved % (this material)",
    "BOE Number", "BOE Date", "Import PO Number", "Qty Imported (this BOE, KGS)",
    "Value Imported / Duty Saved (this BOE, INR)",
]
_MAX_HEADER_SCAN_ROW = 5  # give up looking for the header past this row - same tolerance as rodtep.py


# ── Parsed-row shape ─────────────────────────────────────────────────────────

@dataclass
class ParsedAdvanceLicenseMaterial:
    material_description: str
    itchs_code: str
    qty_authorized: object
    cif_value_authorized: object
    duty_saved_pct: object
    boe_number: str
    boe_date: object
    import_po_number: str
    qty_imported: object
    value_imported: object


@dataclass
class ParsedAdvanceLicense:
    license_number: str
    issue_date: object
    iec: str
    cif_value_authorized: object
    fob_value_export_target: object
    export_product_description: str
    export_validity_date: object
    import_validity_date: object
    status: str
    materials: list = field(default_factory=list)


class HeaderMismatch(Exception):
    pass


def _strip(v):
    return (v or "").strip() if isinstance(v, str) else v


def _find_header_row(ws) -> int:
    """Scans rows 1..._MAX_HEADER_SCAN_ROW for a row whose cell values match
    EXPECTED_HEADERS exactly. Returns that row number. Raises HeaderMismatch
    if none of the scanned rows match."""
    for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=_MAX_HEADER_SCAN_ROW), start=1):
        values = [_strip(cell.value) for cell in row[: len(EXPECTED_HEADERS)]]
        if values == EXPECTED_HEADERS:
            return row_num
    raise HeaderMismatch(
        f"Expected header {EXPECTED_HEADERS} not found in rows 1-{_MAX_HEADER_SCAN_ROW}."
    )


# ── Public entry point ───────────────────────────────────────────────────────

def parse_advance_license_xlsx(file_bytes: bytes) -> list[ParsedAdvanceLicense]:
    """Reads the Advance License workbook (single sheet) and returns one
    ParsedAdvanceLicense per distinct License Number, with its material/usage
    rows grouped underneath in source order. Raises HeaderMismatch if the
    expected header isn't found within the first few rows. A row with a
    blank License Number is skipped (an empty trailing row)."""
    # read_only=True + stream_rows() - see stock.py's/rodtep.py's own
    # comments for why (a real production OOM already hit this app, and
    # ws.max_row/max_column can be None with random access being O(n) per
    # call under read_only - see CLAUDE.md).
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row = _find_header_row(ws)
    data_start_row = header_row + 1
    max_col = len(EXPECTED_HEADERS)

    licenses_by_number: dict[str, ParsedAdvanceLicense] = {}
    license_sequence: list[str] = []

    for r, c in stream_rows(ws, data_start_row, max_col):
        license_number = to_code_str(c[1].value)
        if not license_number:
            continue

        if license_number not in licenses_by_number:
            licenses_by_number[license_number] = ParsedAdvanceLicense(
                license_number=license_number,
                issue_date=to_date(c[2].value),
                iec=to_code_str(c[3].value),
                cif_value_authorized=to_decimal(c[4].value) or 0,
                fob_value_export_target=to_decimal(c[5].value) or 0,
                export_product_description=to_str(c[6].value),
                export_validity_date=to_date(c[7].value),
                import_validity_date=to_date(c[8].value),
                status=to_str(c[9].value),
            )
            license_sequence.append(license_number)

        licenses_by_number[license_number].materials.append(
            ParsedAdvanceLicenseMaterial(
                material_description=to_str(c[10].value),
                itchs_code=to_code_str(c[11].value),
                qty_authorized=to_decimal(c[12].value),
                cif_value_authorized=to_decimal(c[13].value),
                duty_saved_pct=to_decimal(c[14].value),
                boe_number=to_code_str(c[15].value),
                boe_date=to_date(c[16].value),
                import_po_number=to_code_str(c[17].value),
                qty_imported=to_decimal(c[18].value),
                value_imported=to_decimal(c[19].value),
            )
        )

    return [licenses_by_number[num] for num in license_sequence]
