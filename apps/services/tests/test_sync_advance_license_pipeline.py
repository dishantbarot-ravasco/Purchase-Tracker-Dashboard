"""
Integration tests for apps/core/management/commands/sync_advance_license.py -
real Postgres, call_command against a local --file, no mocking (same
convention as test_sync_rodtep_pipeline.py).

Covers what's genuinely different about this command: change detection is a
whole-license hash (like sync_po_csv.py's _po_hash), materials are deleted
and rebuilt together rather than diffed individually, and a license with
multiple usage rows for the same material (e.g. drawn from more than one
BOE) must keep every row, not collapse them.
"""

import io

import openpyxl
import pytest
from django.core.management import call_command

from apps.core.models import AdvanceLicense, SyncRun
from apps.services.parsers.advance_license import EXPECTED_HEADERS


def _build_workbook(license_rows, header_row=1):
    """Builds a minimal real xlsx matching the Advance License ledger shape.
    `license_rows` is a list of (license_number, export_product_description,
    cif_value_authorized, fob_value_export_target, export_validity,
    material_description, boe_number) tuples - one row per material/usage
    instance, same shape as the real source file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for col_idx, header in enumerate(EXPECTED_HEADERS, start=1):
        ws.cell(row=header_row, column=col_idx, value=header)
    for i, row in enumerate(license_rows):
        license_number, export_product_description, cif, fob, export_validity, material_description, boe_number = row
        r = header_row + 1 + i
        ws.cell(row=r, column=1, value=license_number)
        ws.cell(row=r, column=2, value="26/09/2025")
        ws.cell(row=r, column=3, value="0394028678")
        ws.cell(row=r, column=4, value=cif)
        ws.cell(row=r, column=5, value=fob)
        ws.cell(row=r, column=6, value=export_product_description)
        ws.cell(row=r, column=7, value=export_validity)
        ws.cell(row=r, column=8, value="26/09/2026")
        ws.cell(row=r, column=9, value="Active")
        ws.cell(row=r, column=10, value=material_description)
        ws.cell(row=r, column=11, value="40012200")
        ws.cell(row=r, column=12, value=1000)
        ws.cell(row=r, column=13, value=cif)
        ws.cell(row=r, column=14, value=30)
        if boe_number:
            ws.cell(row=r, column=15, value=boe_number)
            ws.cell(row=r, column=16, value="01/06/2026")
            ws.cell(row=r, column=17, value="1000001477")
            ws.cell(row=r, column=18, value=500)
            ws.cell(row=r, column=19, value=100000)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncAdvanceLicense:
    def test_parses_license_and_records_success_syncrun(self, tmp_path):
        fixture_path = tmp_path / "advance_license.xlsx"
        fixture_path.write_bytes(_build_workbook([
            ("0311047672", "Conveyor Belting", 47410000, 53514900, "26/03/2027", "Synthetic Fabric", ""),
        ]))

        call_command("sync_advance_license", file=str(fixture_path))

        assert AdvanceLicense.objects.count() == 1
        lic = AdvanceLicense.objects.get(license_number="0311047672")
        assert lic.export_product_description == "Conveyor Belting"
        assert lic.cif_value_authorized == 47410000
        assert lic.fob_value_export_target == 53514900
        assert lic.materials.count() == 1

        run = SyncRun.objects.filter(plant=SyncRun.Plant.COMPANY, source=SyncRun.Source.ADVANCE_LICENSE).latest("started_at")
        assert run.status == SyncRun.Status.SUCCESS
        assert run.rows_seen == 1
        assert run.rows_changed == 1

    def test_header_on_row_2_after_blank_row_1_still_parses(self, tmp_path):
        fixture_path = tmp_path / "advance_license.xlsx"
        fixture_path.write_bytes(_build_workbook(
            [("0311047672", "Conveyor Belting", 47410000, 53514900, "26/03/2027", "Synthetic Fabric", "")],
            header_row=2,
        ))

        call_command("sync_advance_license", file=str(fixture_path))

        assert AdvanceLicense.objects.count() == 1

    def test_second_run_with_no_changes_updates_nothing(self, tmp_path):
        fixture_path = tmp_path / "advance_license.xlsx"
        fixture_path.write_bytes(_build_workbook([
            ("0311047672", "Conveyor Belting", 47410000, 53514900, "26/03/2027", "Synthetic Fabric", ""),
        ]))

        call_command("sync_advance_license", file=str(fixture_path))
        call_command("sync_advance_license", file=str(fixture_path))

        run = SyncRun.objects.filter(plant=SyncRun.Plant.COMPANY, source=SyncRun.Source.ADVANCE_LICENSE).latest("started_at")
        assert run.rows_changed == 0
        assert AdvanceLicense.objects.count() == 1

    def test_multiple_materials_for_one_license_all_kept(self, tmp_path):
        fixture_path = tmp_path / "advance_license.xlsx"
        fixture_path.write_bytes(_build_workbook([
            ("0311047672", "Conveyor Belting", 47410000, 53514900, "26/03/2027", "Synthetic Fabric", ""),
            ("0311047672", "Conveyor Belting", 47410000, 53514900, "26/03/2027", "Natural Rubber", ""),
        ]))

        call_command("sync_advance_license", file=str(fixture_path))

        lic = AdvanceLicense.objects.get(license_number="0311047672")
        assert lic.materials.count() == 2

    def test_same_material_multiple_usage_rows_all_kept(self, tmp_path):
        """A material drawn against more than one BOE (e.g. 0311051817's own
        SBR material against 3 separate BOEs, see the real filled ledger)
        must keep every usage row, not collapse them into one."""
        fixture_path = tmp_path / "advance_license.xlsx"
        fixture_path.write_bytes(_build_workbook([
            ("0311051817", "Conveyor Belting", 81782250, 92698100, "25/08/2027", "SBR", "3448562"),
            ("0311051817", "Conveyor Belting", 81782250, 92698100, "25/08/2027", "SBR", "2956922"),
        ]))

        call_command("sync_advance_license", file=str(fixture_path))

        lic = AdvanceLicense.objects.get(license_number="0311051817")
        assert lic.materials.count() == 2
        assert {m.boe_number for m in lic.materials.all()} == {"3448562", "2956922"}

    def test_a_license_row_change_rebuilds_its_materials(self, tmp_path):
        fixture_path = tmp_path / "advance_license.xlsx"
        fixture_path.write_bytes(_build_workbook([
            ("0311047672", "Conveyor Belting", 47410000, 53514900, "26/03/2027", "Synthetic Fabric", ""),
        ]))
        call_command("sync_advance_license", file=str(fixture_path))

        fixture_path.write_bytes(_build_workbook([
            ("0311047672", "Conveyor Belting", 47410000, 53514900, "26/03/2027", "Natural Rubber", ""),
        ]))
        call_command("sync_advance_license", file=str(fixture_path))

        lic = AdvanceLicense.objects.get(license_number="0311047672")
        assert lic.materials.count() == 1
        assert lic.materials.first().material_description == "Natural Rubber"

        run = SyncRun.objects.filter(plant=SyncRun.Plant.COMPANY, source=SyncRun.Source.ADVANCE_LICENSE).latest("started_at")
        assert run.rows_changed == 1

    def test_missing_header_records_a_failed_syncrun_and_raises(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "not the real header"
        buf = io.BytesIO()
        wb.save(buf)
        fixture_path = tmp_path / "advance_license.xlsx"
        fixture_path.write_bytes(buf.getvalue())

        with pytest.raises(SystemExit):
            call_command("sync_advance_license", file=str(fixture_path))

        run = SyncRun.objects.filter(plant=SyncRun.Plant.COMPANY, source=SyncRun.Source.ADVANCE_LICENSE).latest("started_at")
        assert run.status == SyncRun.Status.FAILED
