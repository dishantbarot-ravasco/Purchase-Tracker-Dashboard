"""
Integration tests for RTP-Achhad's sync_achhad_po_csv/sync_achhad_mir/
sync_achhad_stock management commands - previously zero coverage. Mirrors
test_sync_po_csv_pipeline.py/test_sync_mir_pipeline.py's HRS conventions
(real Postgres, call_command against a local --file, no mocking), adjusted
for what's genuinely different about Achhad's own file shapes:

- MIR sheet is named "R.M. " (trailing space), header row 2, data from row 3
  (HRS: "RAW MATERIAL", header row 6) - see apps/services/parsers/
  achhad_mir.py's module docstring.
- Stock has no vendor/party-name column at all (one row per material, not
  per (material, vendor) lot like HRS), the sheet's single tab is read
  positionally (wb.sheetnames[0], renamed every month) rather than by a
  fixed name, and row 1 must carry a literal "RM STOCK - DD.MM.YYYY" title
  the parser extracts a year/month from for its daily-movement matrix - see
  apps/services/parsers/achhad_stock.py's module docstring.
"""

import csv
import io

import openpyxl
import pytest
from django.core.management import call_command

from apps.core.models import (
    RTPAchhadDomesticPOLineItem,
    RTPAchhadDomesticPurchaseOrder,
    RTPAchhadMIREntry,
    RTPAchhadRMLot,
    SyncRun,
)
from apps.services.parsers.achhad_mir import DATA_START_ROW as MIR_DATA_START_ROW
from apps.services.parsers.achhad_mir import EXPECTED_HEADERS as MIR_HEADERS
from apps.services.parsers.achhad_mir import HEADER_ROW as MIR_HEADER_ROW
from apps.services.parsers.achhad_mir import SHEET_NAME as MIR_SHEET_NAME
from apps.services.parsers.achhad_stock import DATA_START_ROW as STOCK_DATA_START_ROW
from apps.services.parsers.achhad_stock import EXPECTED_HEADERS as STOCK_HEADERS
from apps.services.parsers.achhad_stock import HEADER_ROW as STOCK_HEADER_ROW
from apps.services.parsers.po_csv import EXPECTED_HEADER as PO_CSV_HEADER

# ── PO CSV fixture ───────────────────────────────────────────────────────────

_PO_ROW = {
    "PO Drive Folder Name": "PO_4000000551", "PO Number": "4000000551",
    "PO Created Date": "01-05-2026", "Vendor Name": "Ganesh Chemicals Ltd",
    "Vendor Address": "GIDC, Achhad", "Vendor GSTIN": "24BBBBB0000B1Z5",
    "Vendor Email": "sales@ganesh.example", "Vendor Code": "V002",
    "Billing Address": "Ravasco, Achhad", "ShipTo": "Ravasco, Achhad",
    "Item Id": "1", "Material Description": "Stearic Acid", "HSN ": "3823",
    "QTY": "500", "UOM": "KG", "Delivery Date ": "15-05-2026",
    "Payment Terms ": "60 Days", "IncoTerms": "DDP", "Currency ": "INR",
    "Net Price": "80.00", "Net Value": "40000.00", "Total Value": "40000.00",
    "Tax Type": "IGST", "Total Inclusive Value": "47200.00", "Remarks": "",
    "PO Number | Item Id": "4000000551|1",
}


def _write_po_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=PO_CSV_HEADER)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncAchhadPoCsv:
    def test_second_run_with_no_changes_updates_nothing(self, tmp_path):
        fixture_path = tmp_path / "po.csv"
        fixture_path.write_text(_write_po_csv([_PO_ROW]), encoding="utf-8")

        call_command("sync_achhad_po_csv", file=str(fixture_path))
        first = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_ACHHAD, source=SyncRun.Source.PO_CSV).latest("started_at")
        assert first.rows_changed == 1
        assert RTPAchhadDomesticPurchaseOrder.objects.get(po_number="4000000551")

        call_command("sync_achhad_po_csv", file=str(fixture_path))
        second = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_ACHHAD, source=SyncRun.Source.PO_CSV).latest("started_at")
        assert second.rows_changed == 0

    def test_a_real_field_change_is_picked_up_on_resync(self, tmp_path):
        fixture_path = tmp_path / "po.csv"
        fixture_path.write_text(_write_po_csv([_PO_ROW]), encoding="utf-8")
        call_command("sync_achhad_po_csv", file=str(fixture_path))

        changed = dict(_PO_ROW, **{"QTY": "600"})
        fixture_path.write_text(_write_po_csv([changed]), encoding="utf-8")
        call_command("sync_achhad_po_csv", file=str(fixture_path))

        item = RTPAchhadDomesticPOLineItem.objects.get(purchase_order__po_number="4000000551")
        assert item.qty == 600


# ── MIR fixture ──────────────────────────────────────────────────────────────

# (party_name, material_description, qty, rate)
_MIR_ROW_A = ("Ganesh Chemicals Ltd", "Stearic Acid", 500, 80.00)
_MIR_ROW_B = ("Kedar Metals Pvt Ltd", "Zinc Oxide", 300, 205.00)


def _write_mir_row(ws, row_idx, party_name, material_description, qty, rate):
    ws[f"A{row_idx}"] = "May-26"
    ws[f"B{row_idx}"] = "AMR001"
    ws[f"C{row_idx}"] = "2026-05-01"
    ws[f"D{row_idx}"] = party_name
    ws[f"E{row_idx}"] = "Gujarat"
    ws[f"F{row_idx}"] = "INV001"
    ws[f"G{row_idx}"] = "2026-05-01"
    ws[f"H{row_idx}"] = "4000000551"
    ws[f"I{row_idx}"] = "2026-04-25"
    ws[f"J{row_idx}"] = material_description
    ws[f"K{row_idx}"] = qty
    ws[f"L{row_idx}"] = "KG"
    ws[f"M{row_idx}"] = rate
    ws[f"N{row_idx}"] = qty * rate
    for col in ("O", "P", "Q", "S", "T", "U", "V", "W", "X", "Y", "AA", "AB"):
        ws[f"{col}{row_idx}"] = 0
    ws[f"R{row_idx}"] = qty * rate
    ws[f"Z{row_idx}"] = qty * rate
    ws[f"AC{row_idx}"] = qty * rate
    ws[f"AD{row_idx}"] = "RTP-ACHHAD"
    ws[f"AE{row_idx}"] = "Production"
    ws[f"AF{row_idx}"] = "Raw Material"
    ws[f"AG{row_idx}"] = ""


def _build_mir_workbook(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = MIR_SHEET_NAME
    for col, header in MIR_HEADERS.items():
        ws[f"{col}{MIR_HEADER_ROW}"] = header
    for offset, row in enumerate(rows):
        _write_mir_row(ws, MIR_DATA_START_ROW + offset, *row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncAchhadMir:
    def test_second_run_with_no_changes_updates_nothing(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        fixture_path.write_bytes(_build_mir_workbook([_MIR_ROW_A, _MIR_ROW_B]))

        call_command("sync_achhad_mir", file=str(fixture_path))
        first = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_ACHHAD, source=SyncRun.Source.MIR).latest("started_at")
        assert first.rows_changed == 2

        call_command("sync_achhad_mir", file=str(fixture_path))
        second = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_ACHHAD, source=SyncRun.Source.MIR).latest("started_at")
        assert second.rows_changed == 0

    def test_a_row_no_longer_in_the_sheet_is_deactivated_not_deleted(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        fixture_path.write_bytes(_build_mir_workbook([_MIR_ROW_A, _MIR_ROW_B]))
        call_command("sync_achhad_mir", file=str(fixture_path))

        row_b_ref = str(MIR_DATA_START_ROW + 1)
        fixture_path.write_bytes(_build_mir_workbook([_MIR_ROW_A]))
        call_command("sync_achhad_mir", file=str(fixture_path))

        entry_b = RTPAchhadMIREntry.objects.get(source_row_ref=row_b_ref)
        assert entry_b.is_active is False

    def test_wrong_sheet_name_records_a_failed_syncrun_and_raises(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Not The Right Sheet"
        buf = io.BytesIO()
        wb.save(buf)
        fixture_path.write_bytes(buf.getvalue())

        with pytest.raises(SystemExit):
            call_command("sync_achhad_mir", file=str(fixture_path))

        run = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_ACHHAD, source=SyncRun.Source.MIR).latest("started_at")
        assert run.status == SyncRun.Status.FAILED
        assert not RTPAchhadMIREntry.objects.exists()


# ── Stock fixture ────────────────────────────────────────────────────────────
# (sap_code, description, rate, opening, received, issued, closing)
_STOCK_ROW_A = ("H1", "SBR 1502", 100.50, 4200, 0, 150, 4050)
_STOCK_ROW_B = ("H2", "Zinc Oxide", 205.00, 1000, 0, 50, 950)


def _write_stock_row(ws, row_idx, sap_code, description, rate, opening, received, issued, closing):
    ws[f"A{row_idx}"] = row_idx  # overall serial (unlabeled column, not header-checked)
    ws[f"B{row_idx}"] = 1
    ws[f"C{row_idx}"] = description
    ws[f"D{row_idx}"] = "2026-05-01"
    ws[f"E{row_idx}"] = sap_code
    ws[f"F{row_idx}"] = rate
    ws[f"G{row_idx}"] = "Zone1"
    ws[f"H{row_idx}"] = 500
    ws[f"I{row_idx}"] = opening
    ws[f"J{row_idx}"] = received
    ws[f"K{row_idx}"] = issued
    ws[f"L{row_idx}"] = closing
    ws[f"M{row_idx}"] = closing * rate
    ws[f"N{row_idx}"] = closing


def _build_stock_workbook(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aug 26-27"
    ws["A1"] = "RM STOCK - 01.05.2026"
    for col, header in STOCK_HEADERS.items():
        ws[f"{col}{STOCK_HEADER_ROW}"] = header
    # A category-divider row (only C filled) right before the data rows -
    # the parser tracks the most recent one seen and stamps it onto rows
    # below (see achhad_stock.py's module docstring).
    divider_row = STOCK_DATA_START_ROW
    ws[f"C{divider_row}"] = "Natural Rubber"
    for offset, row in enumerate(rows):
        _write_stock_row(ws, divider_row + 1 + offset, *row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncAchhadStock:
    def test_second_run_with_no_changes_updates_nothing(self, tmp_path):
        fixture_path = tmp_path / "stock.xlsx"
        fixture_path.write_bytes(_build_stock_workbook([_STOCK_ROW_A, _STOCK_ROW_B]))

        call_command("sync_achhad_stock", file=str(fixture_path), no_snapshot=True)
        first = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_ACHHAD, source=SyncRun.Source.STOCK).latest("started_at")
        assert first.rows_changed == 2
        lot = RTPAchhadRMLot.objects.get(sap_code="H1")
        assert lot.category == "Natural Rubber", "the category-divider row above the data wasn't stamped onto it"

        call_command("sync_achhad_stock", file=str(fixture_path), no_snapshot=True)
        second = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_ACHHAD, source=SyncRun.Source.STOCK).latest("started_at")
        assert second.rows_changed == 0

    def test_no_vendor_field_exists_on_this_plants_stock_lot(self):
        """Achhad's Stock sheet is one row per material, full stop - no
        per-vendor split like HRS's (see RTPAchhadRMLot's own docstring)."""
        assert not hasattr(RTPAchhadRMLot(), "party_name")

    def test_missing_title_date_records_a_failed_syncrun_and_raises(self, tmp_path):
        fixture_path = tmp_path / "stock.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Aug 26-27"
        ws["A1"] = "no date here"
        for col, header in STOCK_HEADERS.items():
            ws[f"{col}{STOCK_HEADER_ROW}"] = header
        buf = io.BytesIO()
        wb.save(buf)
        fixture_path.write_bytes(buf.getvalue())

        with pytest.raises(SystemExit):
            call_command("sync_achhad_stock", file=str(fixture_path))

        run = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_ACHHAD, source=SyncRun.Source.STOCK).latest("started_at")
        assert run.status == SyncRun.Status.FAILED
