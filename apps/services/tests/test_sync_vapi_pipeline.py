"""
Integration tests for RTP-Vapi's sync_vapi_po_csv/sync_vapi_mir/
sync_vapi_stock management commands - previously zero coverage. Mirrors
test_sync_po_csv_pipeline.py/test_sync_mir_pipeline.py's HRS conventions
(real Postgres, call_command against a local --file, no mocking), adjusted
for what's genuinely different about Vapi's own file shapes:

- MIR sheet is named " MIR FILE 26-27 RM" (leading space), header row 1,
  data from row 2, and has no Net/discount columns at all (straight from
  RATE to TAXABLE VALUE) - see apps/services/parsers/vapi_mir.py's module
  docstring. gst_rate_pct is a whole percentage (18.00), not a fraction
  like HRS/Achhad's 0.18.
- Stock sheet is named "Stock" like HRS's (same fixed name, coincidentally),
  but its header sits one row lower in practice at row 6 preceded by a
  5-row document-control title block (not checked by the parser, unlike
  Achhad's title-date requirement) - see apps/services/parsers/
  vapi_stock.py's module docstring. Has a real Supplier Name vendor column
  (basic_rate/supplier_name, not Achhad's vendor-less shape).
"""

import csv
import io
from decimal import Decimal

import openpyxl
import pytest
from django.core.management import call_command

from apps.core.models import (
    RTPVapiDomesticPOLineItem,
    RTPVapiDomesticPurchaseOrder,
    RTPVapiMIREntry,
    RTPVapiRMLot,
    SyncRun,
)
from apps.services.parsers.po_csv import EXPECTED_HEADER as PO_CSV_HEADER
from apps.services.parsers.vapi_mir import DATA_START_ROW as MIR_DATA_START_ROW
from apps.services.parsers.vapi_mir import EXPECTED_HEADERS as MIR_HEADERS
from apps.services.parsers.vapi_mir import HEADER_ROW as MIR_HEADER_ROW
from apps.services.parsers.vapi_mir import SHEET_NAME as MIR_SHEET_NAME
from apps.services.parsers.vapi_stock import DATA_START_ROW as STOCK_DATA_START_ROW
from apps.services.parsers.vapi_stock import EXPECTED_HEADERS as STOCK_HEADERS
from apps.services.parsers.vapi_stock import HEADER_ROW as STOCK_HEADER_ROW
from apps.services.parsers.vapi_stock import SHEET_NAME as STOCK_SHEET_NAME

# ── PO CSV fixture (shared parser/header with HRS/Achhad) ───────────────────

_PO_ROW = {
    "PO Drive Folder Name": "PO_5000000221", "PO Number": "5000000221",
    "PO Created Date": "01-05-2026", "Vendor Name": "Vapi Polymers Pvt Ltd",
    "Vendor Address": "GIDC, Vapi", "Vendor GSTIN": "24CCCCC0000C1Z5",
    "Vendor Email": "sales@vapipoly.example", "Vendor Code": "V003",
    "Billing Address": "Ravasco, Vapi", "ShipTo": "Ravasco, Vapi",
    "Item Id": "1", "Material Description": "NBR 3345", "HSN ": "4002",
    "QTY": "800", "UOM": "KG", "Delivery Date ": "15-05-2026",
    "Payment Terms ": "60 Days", "IncoTerms": "DDP", "Currency ": "INR",
    "Net Price": "150.00", "Net Value": "120000.00", "Total Value": "120000.00",
    "Tax Type": "IGST", "Total Inclusive Value": "141600.00", "Remarks": "",
    "PO Number | Item Id": "5000000221|1",
}


def _write_po_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=PO_CSV_HEADER)
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncVapiPoCsv:
    def test_second_run_with_no_changes_updates_nothing(self, tmp_path):
        fixture_path = tmp_path / "po.csv"
        fixture_path.write_text(_write_po_csv([_PO_ROW]), encoding="utf-8")

        call_command("sync_vapi_po_csv", file=str(fixture_path))
        first = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_VAPI, source=SyncRun.Source.PO_CSV).latest("started_at")
        assert first.rows_changed == 1
        assert RTPVapiDomesticPurchaseOrder.objects.get(po_number="5000000221")

        call_command("sync_vapi_po_csv", file=str(fixture_path))
        second = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_VAPI, source=SyncRun.Source.PO_CSV).latest("started_at")
        assert second.rows_changed == 0

    def test_a_real_field_change_is_picked_up_on_resync(self, tmp_path):
        fixture_path = tmp_path / "po.csv"
        fixture_path.write_text(_write_po_csv([_PO_ROW]), encoding="utf-8")
        call_command("sync_vapi_po_csv", file=str(fixture_path))

        changed = dict(_PO_ROW, **{"QTY": "900"})
        fixture_path.write_text(_write_po_csv([changed]), encoding="utf-8")
        call_command("sync_vapi_po_csv", file=str(fixture_path))

        item = RTPVapiDomesticPOLineItem.objects.get(purchase_order__po_number="5000000221")
        assert item.qty == 900


# ── MIR fixture ──────────────────────────────────────────────────────────────
# (party_name, material_description, qty, rate)
_MIR_ROW_A = ("Vapi Polymers Pvt Ltd", "NBR 3345", 800, 150.00)
_MIR_ROW_B = ("Kedar Metals Pvt Ltd", "Zinc Oxide", 300, 205.00)


def _write_mir_row(ws, row_idx, party_name, material_description, qty, rate, po_number="1000001500"):
    ws[f"A{row_idx}"] = "May-26"
    ws[f"B{row_idx}"] = "MIR01/01"
    ws[f"C{row_idx}"] = "2026-05-01"
    ws[f"D{row_idx}"] = ""  # SAP P.O - ~100% blank on real data (sap_po_number)
    ws[f"E{row_idx}"] = "GRN001"
    ws[f"F{row_idx}"] = ""
    ws[f"G{row_idx}"] = ""
    ws[f"H{row_idx}"] = ""
    ws[f"I{row_idx}"] = party_name
    ws[f"J{row_idx}"] = "Gujarat"
    ws[f"K{row_idx}"] = po_number  # PURCHASE ORDER - added 2026-09-11, the actually-populated PO-number column (po_number_raw)
    ws[f"L{row_idx}"] = "INV001"
    ws[f"M{row_idx}"] = "2026-05-01"
    ws[f"N{row_idx}"] = material_description
    ws[f"O{row_idx}"] = "IC001"
    ws[f"P{row_idx}"] = qty
    ws[f"Q{row_idx}"] = "KG"
    ws[f"R{row_idx}"] = rate
    ws[f"S{row_idx}"] = qty * rate
    ws[f"T{row_idx}"] = 0
    ws[f"U{row_idx}"] = 18.00  # gst_rate_pct - whole percentage, not a fraction
    ws[f"V{row_idx}"] = qty * rate * 0.18
    ws[f"W{row_idx}"] = 0
    ws[f"X{row_idx}"] = 0
    ws[f"Y{row_idx}"] = 0
    ws[f"Z{row_idx}"] = 0
    ws[f"AA{row_idx}"] = qty * rate * 1.18
    ws[f"AB{row_idx}"] = "Raw Material"
    ws[f"AC{row_idx}"] = "2026-05-02"
    ws[f"AD{row_idx}"] = "2026-05-03"


def _build_mir_workbook(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = MIR_SHEET_NAME
    for col, header in MIR_HEADERS.items():
        ws[f"{col}{MIR_HEADER_ROW}"] = header
    for offset, row in enumerate(rows):
        _write_mir_row(ws, MIR_DATA_START_ROW + offset, *row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncVapiMir:
    def test_second_run_with_no_changes_updates_nothing(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        fixture_path.write_bytes(_build_mir_workbook([_MIR_ROW_A, _MIR_ROW_B]))

        call_command("sync_vapi_mir", file=str(fixture_path))
        first = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_VAPI, source=SyncRun.Source.MIR).latest("started_at")
        assert first.rows_changed == 2

        call_command("sync_vapi_mir", file=str(fixture_path))
        second = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_VAPI, source=SyncRun.Source.MIR).latest("started_at")
        assert second.rows_changed == 0

    def test_gst_rate_is_stored_as_a_whole_percentage_not_a_fraction(self, tmp_path):
        """RTPVapiMIREntry.gst_rate_pct is decimal_places=2 specifically
        because Vapi's sheet stores 18% as the literal cell value 18.00, not
        HRS/Achhad's 0.18 fraction - see vapi_mir.py's own inline comment."""
        fixture_path = tmp_path / "mir.xlsx"
        fixture_path.write_bytes(_build_mir_workbook([_MIR_ROW_A]))
        call_command("sync_vapi_mir", file=str(fixture_path))

        entry = RTPVapiMIREntry.objects.get(source_row_ref=str(MIR_DATA_START_ROW))
        assert entry.gst_rate_pct == Decimal("18.00")

    def test_new_purchase_order_column_populates_po_number_raw(self, tmp_path):
        """Header change 2026-09-11: the new 'PURCHASE ORDER' column (K,
        added at the project owner's request for better matching) feeds
        po_number_raw; the original, still-100%-blank 'SAP P.O' column (D)
        is now captured separately as sap_po_number, not conflated with it."""
        fixture_path = tmp_path / "mir.xlsx"
        fixture_path.write_bytes(_build_mir_workbook([_MIR_ROW_A]))
        call_command("sync_vapi_mir", file=str(fixture_path))

        entry = RTPVapiMIREntry.objects.get(source_row_ref=str(MIR_DATA_START_ROW))
        assert entry.po_number_raw == "1000001500"
        assert entry.sap_po_number == ""

    def test_a_row_no_longer_in_the_sheet_is_deactivated_not_deleted(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        fixture_path.write_bytes(_build_mir_workbook([_MIR_ROW_A, _MIR_ROW_B]))
        call_command("sync_vapi_mir", file=str(fixture_path))

        row_b_ref = str(MIR_DATA_START_ROW + 1)
        fixture_path.write_bytes(_build_mir_workbook([_MIR_ROW_A]))
        call_command("sync_vapi_mir", file=str(fixture_path))

        entry_b = RTPVapiMIREntry.objects.get(source_row_ref=row_b_ref)
        assert entry_b.is_active is False

    def test_wrong_sheet_name_records_a_failed_syncrun_and_raises(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Not The Right Sheet"
        buf = io.BytesIO()
        wb.save(buf)
        fixture_path.write_bytes(buf.getvalue())

        with pytest.raises(SystemExit):
            call_command("sync_vapi_mir", file=str(fixture_path))

        run = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_VAPI, source=SyncRun.Source.MIR).latest("started_at")
        assert run.status == SyncRun.Status.FAILED
        assert not RTPVapiMIREntry.objects.exists()


# ── Stock fixture ────────────────────────────────────────────────────────────
# (hsn_code, description, supplier_name, basic_rate)
_STOCK_ROW_A = ("40012200", "NBR 3345", "Vapi Polymers Pvt Ltd", 150.00)
_STOCK_ROW_B = ("28170000", "Zinc Oxide", "Kedar Metals Pvt Ltd", 205.00)


def _write_stock_row(ws, row_idx, hsn_code, description, supplier_name, basic_rate):
    ws[f"A{row_idx}"] = row_idx - STOCK_DATA_START_ROW + 1
    ws[f"B{row_idx}"] = "RTP-1"
    ws[f"C{row_idx}"] = description
    ws[f"D{row_idx}"] = "Category"
    ws[f"E{row_idx}"] = "Sub"
    ws[f"F{row_idx}"] = "KG"
    ws[f"G{row_idx}"] = 4200
    ws[f"H{row_idx}"] = 0
    ws[f"I{row_idx}"] = 150
    ws[f"J{row_idx}"] = 4050
    ws[f"K{row_idx}"] = basic_rate
    ws[f"L{row_idx}"] = 4050 * basic_rate
    ws[f"M{row_idx}"] = "2026-05-01"
    ws[f"N{row_idx}"] = supplier_name
    ws[f"O{row_idx}"] = "RTP-1"
    ws[f"P{row_idx}"] = "Vapi"
    ws[f"Q{row_idx}"] = hsn_code


def _build_stock_workbook(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = STOCK_SHEET_NAME
    for col, header in STOCK_HEADERS.items():
        ws[f"{col}{STOCK_HEADER_ROW}"] = header
    for offset, row in enumerate(rows):
        _write_stock_row(ws, STOCK_DATA_START_ROW + offset, *row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncVapiStock:
    def test_second_run_with_no_changes_updates_nothing(self, tmp_path):
        fixture_path = tmp_path / "stock.xlsx"
        fixture_path.write_bytes(_build_stock_workbook([_STOCK_ROW_A, _STOCK_ROW_B]))

        call_command("sync_vapi_stock", file=str(fixture_path), no_snapshot=True)
        first = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_VAPI, source=SyncRun.Source.STOCK).latest("started_at")
        assert first.rows_changed == 2
        lot = RTPVapiRMLot.objects.get(hsn_code="40012200")
        assert lot.supplier_name == "Vapi Polymers Pvt Ltd"

        call_command("sync_vapi_stock", file=str(fixture_path), no_snapshot=True)
        second = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_VAPI, source=SyncRun.Source.STOCK).latest("started_at")
        assert second.rows_changed == 0

    def test_header_mismatch_records_a_failed_syncrun_and_raises(self, tmp_path):
        fixture_path = tmp_path / "stock.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = STOCK_SHEET_NAME
        wb.active[f"A{STOCK_HEADER_ROW}"] = "Wrong Header"
        buf = io.BytesIO()
        wb.save(buf)
        fixture_path.write_bytes(buf.getvalue())

        with pytest.raises(SystemExit):
            call_command("sync_vapi_stock", file=str(fixture_path))

        run = SyncRun.objects.filter(plant=SyncRun.Plant.RTP_VAPI, source=SyncRun.Source.STOCK).latest("started_at")
        assert run.status == SyncRun.Status.FAILED
