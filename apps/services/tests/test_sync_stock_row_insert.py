"""
Integration test for Snapshot Pipeline Rebuild Phase A - the row-insert
defect this whole phase exists to fix (see apps/services/stock_identity.py
and sync_stock.py's module docstrings): a stock lot used to be keyed on
`source_row_ref` (the raw openpyxl row index), so inserting one row
mid-sheet silently re-labeled an existing lot as whatever material now
occupied its old row, while that lot's HRSStockSnapshot history stayed
attached by foreign key.

**Deliberate, narrow exception to CLAUDE.md's "the Drive-sync/parse/
matching pipeline itself... has no automated test coverage" note.** That
note is still true for the pipeline in general (no test coverage for the
real Drive API calls or the day-to-day sync/match commands as a whole) -
this file is a one-off, specifically because the row-insert fix is
unverifiable any other way: only a real call_command("sync_stock", ...)
run twice against a fixture workbook that's mutated in between proves a
lot's identity survives a row shift. Real Postgres (`@pytest.mark.django_db`,
`uv run pytest`), no mocking - same DB-access convention as
apps/api/tests/test_auth_flow.py.
"""

import io

import openpyxl
import pytest
from django.core.management import call_command

from apps.core.models import HRSStockLot

_HEADER_ROW = 6
_DATA_START_ROW = 7

# Matches apps/services/parsers/stock.py's EXPECTED_HEADERS exactly (A-N are
# header-checked; O/P are genuinely unlabeled in the real file).
_HEADERS = {
    "A": "S.No.", "B": "Description", "C": "SAP ITEM CODE", "D": "Category",
    "E": "Sub Category", "F": "UOM", "G": "Opening\nStock", "H": "REC", "I": "ISSUE",
    "J": "Today\nStock", "K": "Basic Rate", "L": "Value", "M": "Rec. DT.", "N": "No of Days",
}

# (sr_no, description, sap_item_code, party_name) - todays_stock/basic_rate
# are fixed per row below, only identity fields vary here.
_ROW_A = (1, "SBR 1502", "H1", "Rubamin")
_ROW_B = (2, "Zinc Oxide", "H2", "Kedar Metals")
_ROW_C = (3, "Stearic Acid", "H3", "Ganesh")
_INSERTED_ROW = (0, "Carbon Black N330", "H4", "Phillips")


def _write_row(ws, row_idx, sr_no, description, code, vendor):
    ws[f"A{row_idx}"] = sr_no
    ws[f"B{row_idx}"] = description
    ws[f"C{row_idx}"] = code
    ws[f"D{row_idx}"] = "Category"
    ws[f"E{row_idx}"] = "Sub"
    ws[f"F{row_idx}"] = "KG"
    ws[f"G{row_idx}"] = 4200
    ws[f"H{row_idx}"] = 0
    ws[f"I{row_idx}"] = 150
    ws[f"J{row_idx}"] = 4050
    ws[f"K{row_idx}"] = 105.5
    ws[f"L{row_idx}"] = 427275
    ws[f"M{row_idx}"] = None
    ws[f"N{row_idx}"] = 5
    ws[f"O{row_idx}"] = vendor
    ws[f"P{row_idx}"] = "HRS"


def _build_workbook(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock"
    for col, header in _HEADERS.items():
        ws[f"{col}{_HEADER_ROW}"] = header
    for offset, (sr_no, description, code, vendor) in enumerate(rows):
        _write_row(ws, _DATA_START_ROW + offset, sr_no, description, code, vendor)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncStockSurvivesARowInsert:
    def test_lot_identity_and_snapshot_history_survive_a_row_insert(self, tmp_path):
        fixture_path = tmp_path / "stock.xlsx"

        fixture_path.write_bytes(_build_workbook([_ROW_A, _ROW_B, _ROW_C]))
        call_command("sync_stock", file=str(fixture_path))

        zinc_oxide = HRSStockLot.objects.get(sap_item_code="H2")
        original_id = zinc_oxide.id
        original_snapshot_count = zinc_oxide.snapshots.count()
        assert original_snapshot_count == 1

        # Insert a new row at the TOP of the data range - every existing
        # row's sheet position shifts down by one, exactly the incident
        # stock_identity.py's docstring describes.
        fixture_path.write_bytes(_build_workbook([_INSERTED_ROW, _ROW_A, _ROW_B, _ROW_C]))
        call_command("sync_stock", file=str(fixture_path))

        zinc_oxide.refresh_from_db()
        assert zinc_oxide.id == original_id, (
            "Zinc Oxide's lot id changed after a row insert - natural_key isn't actually stable."
        )
        assert zinc_oxide.snapshots.count() == original_snapshot_count, (
            "Snapshot history was duplicated/lost across the row-shifted re-sync."
        )
        assert HRSStockLot.objects.filter(sap_item_code="H4", description="Carbon Black N330").exists()
        assert HRSStockLot.objects.count() == 4
