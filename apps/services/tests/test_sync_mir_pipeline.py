"""
Integration tests for apps/core/management/commands/sync_mir.py -
previously zero coverage (see test_sync_po_csv_pipeline.py's own docstring
for why this file exists and the conventions it follows: real Postgres,
call_command against a local --file, no mocking).

Covers what's genuinely different about this command versus sync_po_csv.py:
per-field sync_utils.unchanged() change detection (not a whole-row hash) and
row deactivation (not deletion) when a row disappears from the sheet - see
sync_mir.py's own module docstring for why source_row_ref isn't a stable
key and is_active exists at all.
"""

import io

import openpyxl
import pytest
from django.core.management import call_command

from apps.core.models import HRSMIREntry, SyncRun
from apps.services.parsers.mir import DATA_START_ROW, EXPECTED_HEADERS, HEADER_ROW, SHEET_NAME

# (party_name, material_description, qty, rate) - everything else fixed below.
_ROW_A = ("Rubamin Private Limited", "SBR 1502", 1000, 100.50)
_ROW_B = ("Kedar Metals Pvt Ltd", "Zinc Oxide", 500, 205.00)


def _write_row(ws, row_idx, party_name, material_description, qty, rate):
    ws[f"A{row_idx}"] = "April-26"
    ws[f"B{row_idx}"] = "MIR001"
    ws[f"C{row_idx}"] = "2026-05-01"
    ws[f"D{row_idx}"] = "GRN001"
    ws[f"E{row_idx}"] = party_name
    ws[f"F{row_idx}"] = "GJ"
    ws[f"G{row_idx}"] = "INV001"
    ws[f"H{row_idx}"] = "2026-05-01"
    ws[f"I{row_idx}"] = ""
    ws[f"J{row_idx}"] = "3000001104"
    ws[f"K{row_idx}"] = "2026-04-25"
    ws[f"L{row_idx}"] = ""
    ws[f"M{row_idx}"] = ""
    ws[f"N{row_idx}"] = material_description
    ws[f"O{row_idx}"] = qty
    ws[f"P{row_idx}"] = "KG"
    ws[f"Q{row_idx}"] = rate
    ws[f"R{row_idx}"] = qty * rate
    for col in ("S", "T", "U", "W", "X", "Y", "Z", "AA", "AB", "AC", "AE", "AF"):
        ws[f"{col}{row_idx}"] = 0
    ws[f"V{row_idx}"] = qty * rate
    ws[f"AD{row_idx}"] = qty * rate
    ws[f"AG{row_idx}"] = qty * rate
    ws[f"AH{row_idx}"] = "HRS"
    ws[f"AI{row_idx}"] = "Production"
    ws[f"AJ{row_idx}"] = "Raw Material"
    ws[f"AK{row_idx}"] = ""


def _build_workbook(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for col, header in EXPECTED_HEADERS.items():
        ws[f"{col}{HEADER_ROW}"] = header
    for offset, row in enumerate(rows):
        _write_row(ws, DATA_START_ROW + offset, *row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncMirIdempotency:
    def test_second_run_with_no_changes_updates_nothing(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        fixture_path.write_bytes(_build_workbook([_ROW_A, _ROW_B]))

        call_command("sync_mir", file=str(fixture_path))
        first_run = SyncRun.objects.filter(source=SyncRun.Source.MIR).latest("started_at")
        assert first_run.rows_changed == 2

        call_command("sync_mir", file=str(fixture_path))
        second_run = SyncRun.objects.filter(source=SyncRun.Source.MIR).latest("started_at")
        assert second_run.rows_seen == 2
        assert second_run.rows_changed == 0, "unchanged MIR rows were re-written on a no-op re-sync"

    def test_a_real_field_change_is_picked_up_on_resync(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        fixture_path.write_bytes(_build_workbook([_ROW_A, _ROW_B]))
        call_command("sync_mir", file=str(fixture_path))

        changed_row_a = ("Rubamin Private Limited", "SBR 1502", 1200, 100.50)  # qty changed
        fixture_path.write_bytes(_build_workbook([changed_row_a, _ROW_B]))
        call_command("sync_mir", file=str(fixture_path))

        run = SyncRun.objects.filter(source=SyncRun.Source.MIR).latest("started_at")
        assert run.rows_changed == 1, "a real qty change on re-sync was treated as a no-op"

        entry = HRSMIREntry.objects.get(source_row_ref=str(DATA_START_ROW))
        assert entry.qty == 1200


@pytest.mark.django_db
class TestSyncMirRowDeactivation:
    def test_a_row_no_longer_in_the_sheet_is_deactivated_not_deleted(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        fixture_path.write_bytes(_build_workbook([_ROW_A, _ROW_B]))
        call_command("sync_mir", file=str(fixture_path))

        row_b_ref = str(DATA_START_ROW + 1)
        assert HRSMIREntry.objects.get(source_row_ref=row_b_ref).is_active is True

        # Row B removed entirely - only row A remains in the sheet.
        fixture_path.write_bytes(_build_workbook([_ROW_A]))
        call_command("sync_mir", file=str(fixture_path))

        entry_b = HRSMIREntry.objects.get(source_row_ref=row_b_ref)
        assert entry_b.is_active is False, "a row removed from the sheet must be deactivated, not left active"
        assert HRSMIREntry.objects.filter(source_row_ref=row_b_ref).exists(), (
            "a row removed from the sheet must be deactivated (soft), never hard-deleted - "
            "see sync_mir.py's module docstring on why source_row_ref isn't a stable key"
        )


@pytest.mark.django_db
class TestSyncMirHeaderMismatch:
    def test_wrong_sheet_name_records_a_failed_syncrun_and_raises(self, tmp_path):
        fixture_path = tmp_path / "mir.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Not The Right Sheet"
        buf = io.BytesIO()
        wb.save(buf)
        fixture_path.write_bytes(buf.getvalue())

        with pytest.raises(SystemExit):
            call_command("sync_mir", file=str(fixture_path))

        run = SyncRun.objects.filter(source=SyncRun.Source.MIR).latest("started_at")
        assert run.status == SyncRun.Status.FAILED
        assert not HRSMIREntry.objects.exists()
