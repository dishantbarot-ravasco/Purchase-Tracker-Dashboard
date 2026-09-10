"""
Integration tests for apps/core/management/commands/sync_rodtep.py - real
Postgres, call_command against a local --file, no mocking (same convention
as test_sync_mir_pipeline.py etc.).

Covers what's genuinely different about this command: change detection
keyed on (script_no, sb_number) rather than sr_no (see
RodtepScrollEntry's own docstring), and idempotent re-sync converging to
0 created/updated - the exact idempotency check CLAUDE.md's own "Change-
detection must compare quantized Decimal values" section requires for any
new sync command.
"""

import io

import openpyxl
import pytest
from django.core.management import call_command

from apps.core.models import RodtepScrollEntry, SyncRun
from apps.services.parsers.rodtep import EXPECTED_HEADERS


def _build_workbook(rows, header_row=1):
    """Builds a minimal real xlsx matching the RoDTEP ledger shape - one
    (script_no, sb_number, sanctioned_amount) tuple per row. `header_row`
    lets tests exercise the real header-position inconsistency this
    parser's own module docstring documents (confirmed against two real
    live files - one has the header on row 1, another on row 2 after a
    blank row 1)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for col_idx, header in enumerate(EXPECTED_HEADERS, start=1):
        ws.cell(row=header_row, column=col_idx, value=header)
    for i, (script_no, sb_number, sanctioned_amount) in enumerate(rows):
        r = header_row + 1 + i
        ws.cell(row=r, column=2, value=script_no)
        ws.cell(row=r, column=3, value="30/03/2026")
        ws.cell(row=r, column=4, value=sb_number)
        ws.cell(row=r, column=5, value="10.02.2026")
        ws.cell(row=r, column=6, value="177246")
        ws.cell(row=r, column=7, value="10.03.2026")
        ws.cell(row=r, column=9, value="INNSA1")
        ws.cell(row=r, column=10, value=sanctioned_amount)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.django_db
class TestSyncRodtep:
    def test_parses_rows_and_records_success_syncrun(self, tmp_path):
        fixture_path = tmp_path / "rodtep.xlsx"
        fixture_path.write_bytes(_build_workbook([("SCRIPT1", "SB1", 100), ("SCRIPT1", "SB2", 200)]))

        call_command("sync_rodtep", file=str(fixture_path))

        assert RodtepScrollEntry.objects.count() == 2
        total = sum(e.sanctioned_amount for e in RodtepScrollEntry.objects.all())
        assert total == 300

        run = SyncRun.objects.filter(plant=SyncRun.Plant.COMPANY, source=SyncRun.Source.RODTEP).latest("started_at")
        assert run.status == SyncRun.Status.SUCCESS
        assert run.rows_seen == 2
        assert run.rows_changed == 2

    def test_header_on_row_2_after_blank_row_1_still_parses(self, tmp_path):
        """The real header-position inconsistency this parser's own module
        docstring documents - confirmed against two real live files."""
        fixture_path = tmp_path / "rodtep.xlsx"
        fixture_path.write_bytes(_build_workbook([("SCRIPT1", "SB1", 100)], header_row=2))

        call_command("sync_rodtep", file=str(fixture_path))

        assert RodtepScrollEntry.objects.count() == 1

    def test_second_run_with_no_changes_updates_nothing(self, tmp_path):
        fixture_path = tmp_path / "rodtep.xlsx"
        fixture_path.write_bytes(_build_workbook([("SCRIPT1", "SB1", 100)]))

        call_command("sync_rodtep", file=str(fixture_path))
        call_command("sync_rodtep", file=str(fixture_path))

        run = SyncRun.objects.filter(plant=SyncRun.Plant.COMPANY, source=SyncRun.Source.RODTEP).latest("started_at")
        assert run.rows_changed == 0
        assert RodtepScrollEntry.objects.count() == 1

    def test_same_script_different_sb_numbers_both_kept(self, tmp_path):
        """Keyed on (script_no, sb_number), not script_no alone - two real
        Shipping Bills under the same script must both survive, not
        collapse onto one row."""
        fixture_path = tmp_path / "rodtep.xlsx"
        fixture_path.write_bytes(_build_workbook([("SCRIPT1", "SB1", 100), ("SCRIPT1", "SB2", 200)]))

        call_command("sync_rodtep", file=str(fixture_path))

        assert RodtepScrollEntry.objects.filter(script_no="SCRIPT1").count() == 2

    def test_missing_header_records_a_failed_syncrun_and_raises(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "not the real header"
        buf = io.BytesIO()
        wb.save(buf)
        fixture_path = tmp_path / "rodtep.xlsx"
        fixture_path.write_bytes(buf.getvalue())

        with pytest.raises(SystemExit):
            call_command("sync_rodtep", file=str(fixture_path))

        run = SyncRun.objects.filter(plant=SyncRun.Plant.COMPANY, source=SyncRun.Source.RODTEP).latest("started_at")
        assert run.status == SyncRun.Status.FAILED

    def test_one_bad_file_does_not_block_the_others(self, monkeypatch):
        """Regression test for a real bug (found and fixed 2026-09-10,
        reported as "the rodtep script sync is not working" right after a
        new file was added to the Drive folder): a single file with a
        HeaderMismatch used to abort the whole command via an unhandled
        raise, so every other already-working script file silently stopped
        updating too. A bad file must now be skipped (PARTIAL status,
        named in error_detail) while every good file still syncs."""
        good_bytes = _build_workbook([("SCRIPT1", "SB1", 100)])
        bad_wb = openpyxl.Workbook()
        bad_wb.active["A1"] = "not the real header"
        bad_buf = io.BytesIO()
        bad_wb.save(bad_buf)
        bad_bytes = bad_buf.getvalue()

        import apps.services.google_client as google_client

        monkeypatch.setattr(
            google_client, "list_files_in_folder",
            lambda parent_id, name_prefix=None: [
                {"id": "good-id", "name": "RODTEP-JNPT-1.xlsx"},
                {"id": "bad-id", "name": "RODTEP-JNPT-2.xlsx"},
            ],
        )
        monkeypatch.setattr(
            google_client, "download_file_bytes",
            lambda file_id: good_bytes if file_id == "good-id" else bad_bytes,
        )

        call_command("sync_rodtep")

        assert RodtepScrollEntry.objects.filter(script_no="SCRIPT1").exists(), "the good file must still sync"
        run = SyncRun.objects.filter(plant=SyncRun.Plant.COMPANY, source=SyncRun.Source.RODTEP).latest("started_at")
        assert run.status == SyncRun.Status.PARTIAL
        assert "RODTEP-JNPT-2.xlsx" in run.error_detail
        assert run.rows_seen == 1
        assert run.rows_changed == 1

